# -*- coding: utf-8 -*-
"""Caricamento e normalizzazione del file Excel."""
import pandas as pd
import json, os
from pathlib import Path
import streamlit as st
from .config import DATA_FILE, SHEET_NAME, COL, PCT_COLS

# Cache JSON con gli URL reali (relativi alla cartella app)
_DATA_DIR        = Path(__file__).parent.parent / "data"
_FONDIDOC_CACHE  = _DATA_DIR / "fondidoc_cache.json"
_QUANTALYS_CACHE = _DATA_DIR / "quantalys_cache.json"


@st.cache_data(show_spinner=False)
def get_last_data_update() -> str:
    """Data ultimo aggiornamento dati fondi.

    Priorità: banner 'Dati di performance aggiornati al: …' in riga 1 del
    file dati (DATA_FILE). Fallback: data di modifica del file. Ritorna '' se
    non ricavabile.
    """
    import datetime, re
    # 1) banner 'Dati di performance aggiornati al: …' nel file dati stesso
    if os.path.exists(DATA_FILE):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(DATA_FILE, read_only=True)
            ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[0]
            a1 = str(ws.cell(1, 1).value or "")
            wb.close()
            m = re.search(r"(\d{2}/\d{2}/\d{4})", a1)
            if m:
                return m.group(1)
        except Exception:
            pass
    # 2) fallback: mtime del file dati effettivo
    try:
        ts = os.path.getmtime(DATA_FILE)
        return datetime.datetime.fromtimestamp(ts).strftime("%d/%m/%Y")
    except Exception:
        return ""


@st.cache_data(show_spinner=False)
def _load_fondidoc_urls() -> dict:
    if not os.path.exists(_FONDIDOC_CACHE):
        return {}
    with open(_FONDIDOC_CACHE, encoding="utf-8") as f:
        return {k: v for k, v in json.load(f).items() if v}


@st.cache_data(show_spinner=False)
def _load_quantalys_urls() -> dict:
    if not os.path.exists(_QUANTALYS_CACHE):
        return {}
    with open(_QUANTALYS_CACHE, encoding="utf-8") as f:
        return {k: v for k, v in json.load(f).items() if v}


def _normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Strip spazi dai nomi colonne e normalizza i tipi."""
    df.columns = [str(c).strip() for c in df.columns]
    # Rinomina colonne con spazi / caratteri speciali
    rename = {
        "ASSET MANAGEMENT HOUSE": "ASSET MANAGEMENT HOUSE",
        "DIVISA FONDO": "DIVISA FONDO",
        "COMMISSIONE DI GESTIONE": "COMMISSIONE DI GESTIONE",
        "COMMISSIONE DI DISTRIBUZIONE": "COMMISSIONE DI DISTRIBUZIONE",
        "COMMENTI": "COMMENTI",
        "% BPS *ANNUI CF": "BPS ANNUI CF",
        "BPS ANNUI CF": "BPS ANNUI CF",
    }
    df = df.rename(columns=rename)
    # Converti percentuali in float
    for col in PCT_COLS:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    # RATING come int
    if COL["rating"] in df.columns:
        df[COL["rating"]] = pd.to_numeric(df[COL["rating"]], errors="coerce")
    return df


def _header_row(path) -> int:
    """0-based indice di riga delle intestazioni. I file 'arricchita' hanno un
    banner 'Dati di performance aggiornati al: …' in riga 1 e le intestazioni
    in riga 2; i file semplici hanno le intestazioni in riga 1."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[0]
        a1 = str(ws.cell(1, 1).value or "")
        wb.close()
        return 1 if a1.strip().lower().startswith("dati di performance") else 0
    except Exception:
        return 0


def _file_hyperlinks(path, header_row: int) -> dict:
    """Legge gli URL reali (hyperlink) di SCHEDA FONDIDOC / SCHEDA QUANTALYS dal
    file, indicizzati per ISIN: {isin: {'fd': url, 'q': url}}."""
    out = {}
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[0]
        hrow = header_row + 1  # openpyxl è 1-based
        hdr = {str(ws.cell(hrow, c).value).strip(): c
               for c in range(1, ws.max_column + 1) if ws.cell(hrow, c).value}
        c_isin = hdr.get("ISIN")
        c_fd   = hdr.get(COL["url_fondidoc"])
        c_q    = hdr.get(COL["url_quantalys"])
        if c_isin:
            for r in range(hrow + 1, ws.max_row + 1):
                isin = ws.cell(r, c_isin).value
                if not isin:
                    continue
                isin = str(isin).strip()
                rec = {}
                if c_fd and ws.cell(r, c_fd).hyperlink and ws.cell(r, c_fd).hyperlink.target:
                    rec["fd"] = ws.cell(r, c_fd).hyperlink.target
                if c_q and ws.cell(r, c_q).hyperlink and ws.cell(r, c_q).hyperlink.target:
                    rec["q"] = ws.cell(r, c_q).hyperlink.target
                if rec:
                    out[isin] = rec
        wb.close()
    except Exception:
        pass
    return out


@st.cache_data(show_spinner="Caricamento fondi...")
def load_data() -> pd.DataFrame:
    _hrow = _header_row(DATA_FILE)
    df = pd.read_excel(DATA_FILE, sheet_name=SHEET_NAME, dtype=str, header=_hrow)
    df = _normalize_columns(df)
    # Converti percentuali
    for col in PCT_COLS:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    if COL["rating"] in df.columns:
        df[COL["rating"]] = pd.to_numeric(df[COL["rating"]], errors="coerce")

    # URL schede: priorità agli hyperlink del file dati, fallback alle cache JSON
    isin_col = COL["isin"]
    if isin_col in df.columns:
        fd_urls, qly_urls = {}, {}
        if os.path.exists(_FONDIDOC_CACHE):
            with open(_FONDIDOC_CACHE, encoding="utf-8") as f:
                fd_urls = {k: v for k, v in json.load(f).items() if v}
        if os.path.exists(_QUANTALYS_CACHE):
            with open(_QUANTALYS_CACHE, encoding="utf-8") as f:
                qly_urls = {k: v for k, v in json.load(f).items() if v}
        file_links = _file_hyperlinks(DATA_FILE, _hrow)

        def _url(isin, kind, cache):
            isin = str(isin).strip()
            rec = file_links.get(isin)
            if rec and rec.get(kind):
                return rec[kind]
            return cache.get(isin)

        keys = df[isin_col].astype(str).str.strip()
        df[COL["url_fondidoc"]]  = keys.map(lambda i: _url(i, "fd", fd_urls))
        df[COL["url_quantalys"]] = keys.map(lambda i: _url(i, "q",  qly_urls))

    return df
