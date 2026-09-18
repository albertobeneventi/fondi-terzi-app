# -*- coding: utf-8 -*-
r"""
scrape_fidax.py — aggiornamento performance via API JSON di FIDA (fidaonline.com).

Sostituisce il vecchio scrape_performance.py (scraping HTML di fondidoc.it, morto
da quando fondidoc.it e' diventato una SPA Angular). I dati arrivano da un endpoint
pubblico senza autenticazione, quindi funziona anche in GitHub Actions (no PC locale):

    POST https://www.fidaonline.com/data/api/extra-data/get-data/fidax/it
    body: {"FidaCodes": [...codici FIDA...], "Fields": [...campi..., "FIDA_CODE"]}

Il codice FIDA di ogni fondo e' nell'hyperlink della colonna SCHEDA FONDIDOC (col 24):
    https://www.fondidoc.it/d/Index/<FIDA_CODE>/<ISIN>_<nome>

Aggiorna il foglio 'tutti quelli trasferibili' di data/tabella_fondi_arricchita_new.xlsx
(header riga 2, dati da riga 3):
    PERF_1Y_EUR    -> col 14  PERF. 1 ANNO   (cumulata)
    PERF_3Y_EUR    -> col 15  PERF. 3 ANNI   (cumulata)
    PERF_YTD_EUR   -> col 16  PERF. YTD
    PERF_Y1Y_EUR   -> col 17  PERF. 2025     (ultimo anno solare chiuso)
    PERF_Y2Y_EUR   -> col 18  PERF. 2024
    PERF_Y3Y_EUR   -> col 19  PERF. 2023
    PERF_Y4Y_EUR   -> col 20  PERF. 2022
    STD_DEV_1Y_EUR -> col 21  VOLATILITA' (1 anno)
E riscrive il banner data in A1. I link FondiDoc/Quantalys non vengono toccati.
"""
import sys, time, re, datetime
from pathlib import Path
import requests
import openpyxl

EXCEL_FILE = Path(__file__).parent / "data" / "tabella_fondi_arricchita_new.xlsx"
SHEET      = "tutti quelli trasferibili"
API_URL    = "https://www.fidaonline.com/data/api/extra-data/get-data/fidax/it"
BATCH      = 150
DELAY      = 1.2

FIELDS = sorted(["PERF_YTD_EUR", "PERF_1Y_EUR", "PERF_3Y_EUR",
                 "PERF_Y1Y_EUR", "PERF_Y2Y_EUR", "PERF_Y3Y_EUR", "PERF_Y4Y_EUR",
                 "STD_DEV_1Y_EUR", "FIDA_CODE"])

# col Excel -> campo API
COLMAP = {14: "PERF_1Y_EUR", 15: "PERF_3Y_EUR", 16: "PERF_YTD_EUR",
          17: "PERF_Y1Y_EUR", 18: "PERF_Y2Y_EUR", 19: "PERF_Y3Y_EUR", 20: "PERF_Y4Y_EUR",
          21: "STD_DEV_1Y_EUR"}

URL_RE = re.compile(r"fondidoc\.it/d/\w+/([A-Za-z0-9]+)/([A-Z0-9]+)_")


def log(msg):
    print(f"[{datetime.datetime.now():%H:%M:%S}] {msg}", flush=True)


def extract_funds(ws):
    """[(row, isin, fida_code)] dalle righe con URL FondiDoc valido (hyperlink col 24)."""
    out, nourl = [], 0
    for r in range(3, ws.max_row + 1):
        isin = ws.cell(r, 3).value
        if not isin:
            continue
        isin = str(isin).strip()
        fd = ws.cell(r, 24)
        for cand in (fd.hyperlink.target if fd.hyperlink else "",
                     str(fd.value or ""), str(ws.cell(r, 4).value or "")):
            m = URL_RE.search(cand)
            if m:
                out.append((r, isin, m.group(1)))
                break
        else:
            nourl += 1
    return out, nourl


def fetch_batch(session, codes):
    last = None
    for _ in range(3):
        try:
            r = session.post(API_URL, json={"FidaCodes": codes, "Fields": FIELDS}, timeout=40)
            if r.status_code == 200:
                return {d["FIDA_CODE"]: d for d in r.json() if d.get("FIDA_CODE")}
            last = f"HTTP {r.status_code}"
        except Exception as e:
            last = str(e)
        time.sleep(4)
    raise RuntimeError(f"batch fallito: {last}")


def main():
    log("=== AGGIORNAMENTO PERFORMANCE (FIDA fidax API) ===")
    today = datetime.date.today()
    wb = openpyxl.load_workbook(str(EXCEL_FILE))
    ws = wb[SHEET]

    funds, nourl = extract_funds(ws)
    codes = sorted({f[2] for f in funds})
    log(f"righe con codice FIDA: {len(funds)} | codici unici: {len(codes)} | righe senza URL: {nourl}")

    session = requests.Session()
    session.headers.update({
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)",
        "Content-Type": "application/json",
        "Origin": "https://www.fidaonline.com",
        "Referer": "https://www.fidaonline.com/",
    })

    data, got, err = {}, 0, 0
    for i in range(0, len(codes), BATCH):
        chunk = codes[i:i + BATCH]
        try:
            res = fetch_batch(session, chunk)
            data.update(res)
            got += len(res)
        except Exception as e:
            err += len(chunk)
            log(f"  batch {i//BATCH+1}: ERRORE {e}")
        if (i // BATCH) % 10 == 0:
            log(f"  batch {i//BATCH+1}/{(len(codes)+BATCH-1)//BATCH} — trovati {got}, err {err}")
        time.sleep(DELAY)
    log(f"scraping finito: dati {got}, errori batch {err}")

    def num(v):
        return v if isinstance(v, (int, float)) else None

    rows_upd = cells_upd = 0
    for (r, isin, fida) in funds:
        d = data.get(fida)
        if not d:
            continue
        touched = False
        for col, key in COLMAP.items():
            v = num(d.get(key))
            if v is None:
                continue
            if col == 21 and v == 0:      # volatilita 0 = dato assente
                continue
            c = ws.cell(r, col)
            c.value = round(v, 6)
            c.number_format = "0.00%"
            cells_upd += 1
            touched = True
        if touched:
            rows_upd += 1

    ws.cell(1, 1).value = (f"Dati di performance aggiornati al: {today.strftime('%d/%m/%Y')} "
                           f"(fonte: fondidoc.it / FIDA - valuta EUR)")
    wb.save(str(EXCEL_FILE))
    log(f"Excel salvato. righe aggiornate: {rows_upd}, celle: {cells_upd}")
    return 0 if err == 0 else 10


if __name__ == "__main__":
    sys.exit(main() or 0)
